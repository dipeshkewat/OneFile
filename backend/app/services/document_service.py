from pathlib import Path

import fitz
from docx import Document


def docx_to_pdf(source: Path, destination: Path) -> None:
    document = Document(str(source))
    pdf = fitz.open()
    try:
        page = pdf.new_page()
        y = 54.0
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            lines = [text[index : index + 95] for index in range(0, len(text), 95)] or [""]
            for line in lines:
                if y > page.rect.height - 54:
                    page = pdf.new_page()
                    y = 54.0
                if line:
                    page.insert_text((54, y), line, fontsize=11, fontname="helv")
                y += 16
            y += 6
        pdf.save(destination)
    finally:
        pdf.close()


def pdf_to_docx(source: Path, destination: Path) -> None:
    pdf = fitz.open(source)
    document = Document()
    try:
        for page in pdf:
            text = page.get_text("text").strip()
            if text:
                for paragraph in text.splitlines():
                    if paragraph.strip():
                        document.add_paragraph(paragraph.strip())
            else:
                document.add_paragraph("")
        document.save(str(destination))
    finally:
        pdf.close()
